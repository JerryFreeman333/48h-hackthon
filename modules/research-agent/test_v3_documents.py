"""Controlled synthetic office/image fixtures; never claim real company source success."""
import base64
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
import v3_documents as documents
import worker_v3
from v2.transport import SourceError

LEGAL = '浙江大华技术股份有限公司'
IDENTITY = {'company_id': 271, 'legal_name': LEGAL, 'brand': '大华股份',
            'aliases': [LEGAL, '大华股份'], 'credit_code': None, 'match_status': 'record_clue'}

def chinese_image():
    fonts = [Path('C:/Windows/Fonts/msyh.ttc'), Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')]
    font = next((p for p in fonts if p.exists()), None)
    if font is None:
        raise RuntimeError('Chinese controlled OCR validation requires an installed CJK font')
    image = Image.new('RGB', (1500, 500), 'white')
    draw = ImageDraw.Draw(image)
    face = ImageFont.truetype(str(font), 46)
    for index, line in enumerate([LEGAL, '嵌入式软件工程师', '税前固定月薪 18000 元，每周双休。', '合成测试材料，不是真实招聘承诺。']):
        draw.text((60, 50 + 100 * index), line, font=face, fill='black')
    buffer = io.BytesIO(); image.save(buffer, 'PNG')
    return image, buffer.getvalue()

class DocumentTests(unittest.TestCase):
    def test_docx_preserves_paragraph_and_table_cell_coordinates_without_pages(self):
        doc = Document(); doc.add_paragraph(LEGAL)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '岗位'; table.cell(0, 1).text = '嵌入式软件工程师'
        table.cell(1, 0).text = '税前固定月薪'; table.cell(1, 1).text = '18000 元'
        buffer = io.BytesIO(); doc.save(buffer)
        parsed = documents.parse_local_document(buffer.getvalue(), 'docx', '合成 DOCX')
        self.assertIn('税前固定月薪 | 18000 元', parsed['text'])
        self.assertTrue(any(p['locator'].get('column') == 2 and p['locator'].get('row') == 2 for p in parsed['paragraphs']))
        self.assertFalse(parsed['pages']); self.assertFalse(parsed['review_required'])

    def test_excel_preserves_sheet_cell_and_formula_without_evaluation(self):
        book = Workbook(); sheet = book.active; sheet.title = '工资口径'
        sheet.append([LEGAL, '岗位', '固定月薪']); sheet.append(['杭州', '嵌入式软件工程师', 18000])
        sheet['D2'] = '=C2*12'; buffer = io.BytesIO(); book.save(buffer)
        parsed = documents.parse_local_document(buffer.getvalue(), 'xlsx', '合成 XLSX')
        selected = next(p for p in parsed['paragraphs'] if p['locator'].get('cell') == 'D2')
        self.assertEqual(selected['text'], '[未计算公式] =C2*12')
        self.assertEqual(selected['locator']['sheet'], '工资口径')
        self.assertIn('formulas_preserved_not_evaluated', parsed['warnings'])
        self.assertNotIn('216000', parsed['text'])

    def test_csv_quoted_chinese_cells_coordinates_and_literal_formula(self):
        body = f'{LEGAL},岗位,薪酬\n杭州,"嵌入式软件工程师,高级",=18000*12\n'.encode('gb18030')
        parsed = documents.parse_local_document(body, 'csv', '合成 CSV')
        cell = next(p for p in parsed['paragraphs'] if p['locator'].get('row') == 2 and p['locator'].get('column') == 2)
        self.assertEqual(cell['text'], '嵌入式软件工程师,高级')
        self.assertIn('formula_like_text_preserved_not_evaluated', parsed['warnings'])

    def test_wrong_office_kind_macro_external_entity_and_zip_bomb_rejected(self):
        for name, payload, required, expected in [
            ('word/vbaProject.bin', b'x', 'word/document.xml', 'embedded_or_macro_document'),
            ('word/document.xml', b'<!DOCTYPE root [<!ENTITY x SYSTEM "file:///secret">]>', 'word/document.xml', 'xml_external_entities_forbidden'),
            ('word/document.xml', b'a' * 100_000, 'word/document.xml', 'office_zip_expansion_limit'),
        ]:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('[Content_Types].xml', b'<Types/>')
                archive.writestr(required, payload if name == required else b'<doc/>')
                if name != required: archive.writestr(name, payload)
            with self.assertRaises(SourceError) as caught:
                documents.checked_zip(buffer.getvalue(), 'docx')
            self.assertEqual(caught.exception.reason, expected)
        book = Workbook(); buffer = io.BytesIO(); book.save(buffer)
        with self.assertRaises(SourceError) as caught: documents.checked_zip(buffer.getvalue(), 'docx')
        self.assertEqual(caught.exception.reason, 'wrong_office_document_kind')

    def test_local_subprocess_deadline_is_hard_failure(self):
        with patch.object(documents.subprocess, 'run', side_effect=subprocess.TimeoutExpired('parser', 1)):
            with self.assertRaises(SourceError) as caught:
                documents.parse_local_document(b'a,b', 'csv', 'timeout', seconds=1)
        self.assertEqual(caught.exception.status, 'timeout')

    def test_missing_local_ocr_is_explicit_and_never_claims_execution(self):
        with patch.dict('sys.modules', {'rapidocr_onnxruntime': None}):
            with self.assertRaises(SourceError) as caught: documents.ocr_engine()
        self.assertEqual(caught.exception.status, 'not_configured')
        self.assertEqual(caught.exception.reason, 'local_ocr_not_installed')

    def test_real_local_chinese_ocr_image_and_scan_pdf_keep_geometry_and_require_review(self):
        image, png = chinese_image()
        pdf = io.BytesIO(); image.save(pdf, 'PDF', resolution=72.0)
        for kind, body in [('image', png), ('scan_pdf', pdf.getvalue())]:
            parsed = documents.parse_local_document(body, kind, '合成中文 OCR', seconds=30)
            self.assertIn(LEGAL, parsed['text'])
            self.assertIn('18000', parsed['text'])
            self.assertTrue(parsed['review_required'])
            recognized = [p for p in parsed['paragraphs'] if 'bbox' in p['locator']]
            self.assertTrue(recognized)
            self.assertTrue(all(p['locator']['physical_page'] == 1 and len(p['locator']['bbox']) == 8 for p in recognized))
            self.assertTrue(all(0 <= p['locator']['confidence'] <= 1 for p in recognized))
        image.close()

    def test_ocr_page_and_image_limits_enforced(self):
        image = Image.new('RGB', (100, 100), 'white')
        buffer = io.BytesIO(); image.save(buffer, 'PDF', save_all=True, append_images=[image] * 12)
        with self.assertRaises(SourceError) as caught:
            documents.parse_local_document(buffer.getvalue(), 'scan_pdf', 'too many pages')
        self.assertEqual(caught.exception.reason, 'ocr_page_limit')
        with self.assertRaises(SourceError): documents.image_extension(b'GIF89a')
        image.close()

    def test_actual_worker_import_saves_binary_original_and_ocr_review_checkpoint(self):
        class Fetcher:
            def __init__(self, **kwargs): self.requests = 0; self.max_requests = kwargs['max_requests']
            def remaining(self): return 30
        class Search:
            def __init__(self, fetcher): pass
        image, png = chinese_image(); image.close()
        with tempfile.TemporaryDirectory() as folder, patch.object(worker_v3, 'ROOT', Path(folder)), patch.object(worker_v3, 'company_identity', return_value=IDENTITY), patch.object(worker_v3, 'PublicFetcher', Fetcher), patch.object(worker_v3, 'SearchAdapter', Search):
            request = {'company_id': 271, 'question_id': 'q-ocr-r0', 'task_id': 'task-ocr', 'query': '', 'urls': [],
                       'imports': [{'kind': 'image', 'content': base64.b64encode(png).decode(), 'title': '合成截图', 'synthetic_fixture': True}], 'max_requests': 0, 'seconds': 35}
            state = worker_v3.handle(request)
            self.assertEqual(state['requests'], 0)
            failure = [{'mode': r['acquisitionMode'], 'access': r['accessState'], 'analysis': r['analysisState'],
                        'reason': r['failureReason']} for r in state['records']]
            self.assertEqual(len(state['documents']), 1, msg=json.dumps(failure, ensure_ascii=False))
            record = state['records'][0]; document = state['documents'][0]
            self.assertEqual(record['sourceClass'], 'synthetic_fixture')
            self.assertTrue(record['reviewRequired']); self.assertTrue(document['reviewRequired'])
            self.assertTrue(record['rawRef'].endswith('.png'))
            self.assertEqual((Path(folder) / record['rawRef']).read_bytes(), png)
            self.assertTrue(any('bbox' in u['locator'] for u in document['units']))
            with patch.object(worker_v3, 'parse_local_document', side_effect=AssertionError('repeat OCR')):
                self.assertEqual(worker_v3.handle(request), state)

    def test_worker_ocr_deadline_preserves_original_and_does_not_repeat_or_claim_success(self):
        class Fetcher:
            def __init__(self, **kwargs): self.requests = 0; self.max_requests = kwargs['max_requests']
            def remaining(self): return 30
        class Search:
            def __init__(self, fetcher): pass
        image, png = chinese_image(); image.close()
        with tempfile.TemporaryDirectory() as folder, patch.object(worker_v3, 'ROOT', Path(folder)), patch.object(worker_v3, 'company_identity', return_value=IDENTITY), patch.object(worker_v3, 'PublicFetcher', Fetcher), patch.object(worker_v3, 'SearchAdapter', Search):
            request = {'company_id': 271, 'question_id': 'q-timeout-r0', 'task_id': 'task-timeout', 'query': '', 'urls': [],
                       'imports': [{'kind': 'image', 'content': base64.b64encode(png).decode(), 'title': '合成截图', 'synthetic_fixture': True}], 'max_requests': 0, 'seconds': 35}
            with patch.object(worker_v3, 'parse_local_document', side_effect=SourceError('timeout', 'image', 'local_parse_deadline')) as parser:
                state = worker_v3.handle(request)
                self.assertEqual(state['documents'], [])
                self.assertEqual(state['records'][0]['accessState'], 'timeout')
                self.assertEqual(state['records'][0]['failureReason'], 'local_parse_deadline')
                self.assertEqual(state['records'][0]['analysisState'], 'rejected')
                self.assertEqual((Path(folder) / state['records'][0]['rawRef']).read_bytes(), png)
                self.assertEqual(worker_v3.handle(request), state)
                self.assertEqual(parser.call_count, 1)
                self.assertEqual(state['requests'], 0)

if __name__ == '__main__': unittest.main()

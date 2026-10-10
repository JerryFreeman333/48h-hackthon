"""Bounded local, actively supplied documents. No macros, formula evaluation or networking.

OCR is a lossy transcription of a source claim: every OCR-derived unit requires review.
The subprocess deadline includes model loading and kills native inference on timeout.
"""
from __future__ import annotations
import base64
import contextlib
import csv
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import time
import zipfile
from v2.transport import SourceError

VERSION = 'v3-local-documents/1'
MAX_BYTES = 18_000_000
MAX_TEXT = 1_500_000
MAX_CELLS = 25_000
MAX_OCR_PAGES = 12
MAX_IMAGE_PIXELS = 16_000_000

def checked_zip(body: bytes, kind: str):
    """Validate the entire OOXML archive without extracting files to disk."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(body))
        files = archive.infolist()
        if len(files) > 1200 or sum(f.file_size for f in files) > 60_000_000:
            raise SourceError('unsupported', kind, 'office_zip_budget')
        required = 'word/document.xml' if kind == 'docx' else 'xl/workbook.xml'
        names = {f.filename for f in files}
        if required not in names or '[Content_Types].xml' not in names:
            raise SourceError('parse_error', kind, 'wrong_office_document_kind')
        for item in files:
            name = item.filename.replace('\\', '/')
            if name.startswith('/') or '..' in name.split('/') or item.flag_bits & 1:
                raise SourceError('unsupported', kind, 'unsafe_or_encrypted_office_zip')
            if item.file_size > 15_000_000 or item.file_size > max(1, item.compress_size) * 200:
                raise SourceError('unsupported', kind, 'office_zip_expansion_limit')
            if re.search(r'vbaProject|macrosheets|activeX|embeddings/', name, re.I):
                raise SourceError('unsupported', kind, 'embedded_or_macro_document')
            if name.endswith(('.xml', '.rels')):
                data = archive.read(item)
                if re.search(br'<!DOCTYPE|<!ENTITY', data.replace(b'\x00', b''), re.I):
                    raise SourceError('unsupported', kind, 'xml_external_entities_forbidden')
                if name == '[Content_Types].xml' and b'macroEnabled' in data:
                    raise SourceError('unsupported', kind, 'embedded_or_macro_document')
        archive.close()
    except SourceError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise SourceError('parse_error', kind, 'invalid_office_zip')

def result(title, paragraphs, tables=None, pages=None, warnings=None, review=False):
    text = '\n'.join(p['text'] for p in paragraphs)
    if len(text) > MAX_TEXT:
        raise SourceError('unsupported', 'document', 'document_text_limit')
    if not text.strip():
        raise SourceError('parse_error', 'document', 'empty_document')
    return {'title': title[:1000], 'text': text, 'paragraphs': paragraphs,
            'pages': pages or [], 'tables': tables or [], 'published_at': None,
            'links': [], 'origin_url': None, 'parser_version': VERSION,
            'warnings': warnings or [], 'review_required': review}

def paragraph(items, text, **locator):
    value = str(text or '').replace('\x00', '').strip()
    if value:
        items.append({'paragraph': len(items) + 1, 'text': value,
                      'locator': {'paragraph': len(items) + 1, **locator}})
    if len(items) > MAX_CELLS or sum(len(p['text']) for p in items[-1:]) > 100_000:
        raise SourceError('unsupported', 'document', 'document_unit_limit')

def parse_docx(body, title):
    checked_zip(body, 'docx')
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(io.BytesIO(body))
    items, tables, table_index = [], [], 0
    for block in document.element.body.iterchildren():
        if block.tag.endswith('}p'):
            paragraph(items, Paragraph(block, document).text)
        elif block.tag.endswith('}tbl'):
            table_index += 1
            table = Table(block, document)
            rows = []
            for row_index, row in enumerate(table.rows, 1):
                cells = [c.text.strip() for c in row.cells]
                if sum(len(r) for r in rows) + len(cells) > MAX_CELLS:
                    raise SourceError('unsupported', 'docx', 'table_cell_limit')
                rows.append(cells)
                # Full row text remains searchable and exact; table cell coordinates are preserved.
                paragraph(items, ' | '.join(cells), table=table_index, row=row_index)
                for column_index, cell in enumerate(cells, 1):
                    paragraph(items, cell, table=table_index, row=row_index, column=column_index)
            tables.append({'table': table_index, 'rows': rows, 'physical_page': None})
    return result(title, items, tables, warnings=['office_links_not_followed', 'docx_layout_pages_not_inferred',
                                                'headers_footers_comments_embedded_images_not_parsed'])

def scalar(cell):
    if cell.value is None:
        return ''
    # Store formula source, never cached output or an evaluated result.
    if cell.data_type == 'f':
        return '[未计算公式] ' + str(cell.value)
    value = cell.value
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return '[非有限数值]'
    return str(value)

def parse_xlsx(body, title):
    checked_zip(body, 'xlsx')
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
    items, tables, cell_count, formula = [], [], 0, False
    try:
        if len(book.worksheets) > 24:
            raise SourceError('unsupported', 'xlsx', 'sheet_limit')
        for table_index, sheet in enumerate(book.worksheets, 1):
            if sheet.max_row > 5000 or sheet.max_column > 128:
                raise SourceError('unsupported', 'xlsx', 'worksheet_dimension_limit')
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                cell_count += len(row)
                if cell_count > MAX_CELLS:
                    raise SourceError('unsupported', 'xlsx', 'table_cell_limit')
                values = [scalar(c) for c in row]
                if not any(values):
                    continue
                rows.append({'row': row_index, 'values': values})
                paragraph(items, ' | '.join(values), sheet=sheet.title, table=table_index, row=row_index)
                for column_index, (cell, value) in enumerate(zip(row, values), 1):
                    if not value:
                        continue
                    formula |= cell.data_type == 'f'
                    paragraph(items, value, sheet=sheet.title, table=table_index, row=row_index,
                              column=column_index, cell=cell.coordinate)
            tables.append({'table': table_index, 'sheet': sheet.title, 'rows': rows, 'physical_page': None})
    finally:
        book.close()
    warnings = ['office_links_not_followed', 'cell_display_formats_not_inferred']
    if formula:
        warnings.append('formulas_preserved_not_evaluated')
    return result(title, items, tables, warnings=warnings)

def parse_csv(body, title):
    try:
        text = body.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            text = body.decode('gb18030')
        except UnicodeDecodeError:
            raise SourceError('parse_error', 'csv', 'unsupported_csv_encoding')
    if '\x00' in text:
        raise SourceError('parse_error', 'csv', 'binary_csv')
    csv.field_size_limit(100_000)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    items, rows, count, formula = [], [], 0, False
    try:
        for row_index, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
            count += len(row)
            if row_index > 5000 or len(row) > 128 or count > MAX_CELLS:
                raise SourceError('unsupported', 'csv', 'table_cell_limit')
            values = [v.strip() for v in row]
            rows.append(values)
            paragraph(items, ' | '.join(values), table=1, row=row_index)
            for column_index, value in enumerate(values, 1):
                formula |= value.startswith(('=', '+', '-', '@'))
                paragraph(items, value, table=1, row=row_index, column=column_index)
    except csv.Error:
        raise SourceError('parse_error', 'csv', 'csv_parse_failure')
    warnings = ['csv_cells_are_unverified_literal_text']
    if formula:
        warnings.append('formula_like_text_preserved_not_evaluated')
    return result(title, items, [{'table': 1, 'rows': rows, 'physical_page': None}], warnings=warnings)

def image_extension(body):
    if body.startswith(b'\x89PNG\r\n\x1a\n'):
        return '.png'
    if body.startswith(b'\xff\xd8\xff'):
        return '.jpg'
    raise SourceError('unsupported', 'image', 'only_png_jpeg_supported')

def checked_image(body):
    image_extension(body)
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(io.BytesIO(body)) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS or min(image.size) < 16 or getattr(image, 'n_frames', 1) != 1:
                raise SourceError('unsupported', 'image', 'image_dimension_limit')
            image.load()
            return image.convert('RGB')
    except SourceError:
        raise
    except Exception:
        raise SourceError('parse_error', 'image', 'invalid_or_oversized_image')

def ocr_engine():
    try:
        import rapidocr_onnxruntime
        models = Path(rapidocr_onnxruntime.__file__).parent / 'models'
        names = {'det_model_path': 'ch_PP-OCRv4_det_infer.onnx',
                 'cls_model_path': 'ch_ppocr_mobile_v2.0_cls_infer.onnx',
                 'rec_model_path': 'ch_PP-OCRv4_rec_infer.onnx'}
        # Explicit installed local paths: do not allow runtime model downloads.
        paths = {key: str(models / name) for key, name in names.items()}
        if not all(Path(p).is_file() for p in paths.values()):
            raise SourceError('not_configured', 'ocr', 'local_ocr_models_missing')
        return rapidocr_onnxruntime.RapidOCR(**paths, intra_op_num_threads=1,
                                           inter_op_num_threads=1, det_use_cuda=False,
                                           cls_use_cuda=False, rec_use_cuda=False)
    except SourceError:
        raise
    except ImportError:
        raise SourceError('not_configured', 'ocr', 'local_ocr_not_installed')
    except Exception:
        raise SourceError('parse_error', 'ocr', 'local_ocr_initialization_failed')

def recognize(engine, image, page, items):
    import numpy as np
    try:
        detected, _elapsed = engine(np.array(image))
    except Exception:
        raise SourceError('parse_error', 'ocr', 'local_ocr_inference_failed')
    for polygon, text, confidence in detected or []:
        if len(items) >= 2500:
            raise SourceError('unsupported', 'ocr', 'ocr_line_limit')
        bbox = [float(v) for point in polygon for v in point]
        if len(bbox) != 8 or not all(math.isfinite(v) for v in bbox):
            continue
        paragraph(items, text, physical_page=page, bbox=bbox,
                  confidence=min(1.0, max(0.0, float(confidence))), engine='rapidocr-onnxruntime')

def parse_image(body, title):
    image = checked_image(body)
    items = []
    recognize(ocr_engine(), image, 1, items)
    image.close()
    page = {'physical_page': 1, 'printed_page': None, 'text': '\n'.join(p['text'] for p in items), 'ocr': True}
    return result(title, items, pages=[page], warnings=['ocr_transcription_requires_review', 'ocr_bbox_image_pixels'], review=True)

def parse_scan_pdf(body, title):
    import pypdfium2
    if not body.startswith(b'%PDF-'):
        raise SourceError('parse_error', 'pdf', 'invalid_pdf_signature')
    try:
        reader = pypdfium2.PdfDocument(body)
        if len(reader) > MAX_OCR_PAGES:
            reader.close()
            raise SourceError('unsupported', 'ocr', 'ocr_page_limit')
        items, pages, engine = [], [], None
        for index in range(len(reader)):
            page = reader[index]
            if page.get_width() * page.get_height() * 4 > MAX_IMAGE_PIXELS:
                page.close()
                raise SourceError('unsupported', 'ocr', 'render_pixel_limit')
            textpage = page.get_textpage()
            native = textpage.get_text_bounded() or ''
            textpage.close()
            before = len(items)
            is_ocr = len(native.strip()) < 30
            if is_ocr:
                if engine is None:
                    engine = ocr_engine()
                bitmap = page.render(scale=2)
                image = bitmap.to_pil().convert('RGB')
                recognize(engine, image, index + 1, items)
                image.close()
                bitmap.close()
            else:
                for line in native.splitlines():
                    paragraph(items, line, physical_page=index + 1)
            page.close()
            pages.append({'physical_page': index + 1, 'printed_page': None,
                          'text': '\n'.join(p['text'] for p in items[before:]), 'ocr': is_ocr})
        reader.close()
        return result(title, items, pages=pages, warnings=['ocr_transcription_requires_review', 'ocr_bbox_render_pixels_scale_2'], review=True)
    except SourceError:
        raise
    except Exception:
        raise SourceError('parse_error', 'pdf', 'scan_pdf_parse_failure')

def parse_local_document(body: bytes, kind: str, title: str, *, seconds=25):
    """The local child process has no source-fetching code and accepts bytes, never paths."""
    if not isinstance(body, bytes) or len(body) > MAX_BYTES:
        raise SourceError('unsupported', 'import', 'import_size_limit')
    if kind not in ('docx', 'xlsx', 'csv', 'image', 'scan_pdf'):
        raise SourceError('unsupported', 'import', 'unsupported_import_kind')
    payload = json.dumps({'body': base64.b64encode(body).decode('ascii'), 'kind': kind, 'title': title})
    allowed_env = {'PATH', 'SYSTEMROOT', 'WINDIR', 'TEMP', 'TMP', 'TMPDIR', 'LANG', 'LC_ALL'}
    env = {k: v for k, v in os.environ.items() if k.upper() in allowed_env}
    env.update(OMP_NUM_THREADS='1', OPENBLAS_NUM_THREADS='1')
    try:
        completed = subprocess.run([sys.executable, '-X', 'utf8', str(Path(__file__).resolve()), '--parse'],
                                   input=payload, text=True, encoding='utf-8', capture_output=True,
                                   timeout=max(1, min(float(seconds), 35)), env=env)
        data = json.loads(completed.stdout)
        if 'error' in data:
            raise SourceError(data['error']['status'], data['error']['stage'], data['error']['reason'])
        if completed.returncode != 0:
            raise SourceError('parse_error', kind, 'local_parser_failed')
        return data
    except subprocess.TimeoutExpired:
        raise SourceError('timeout', kind, 'local_parse_deadline')
    except SourceError:
        raise
    except Exception:
        raise SourceError('parse_error', kind, 'local_parser_failed')

if __name__ == '__main__':
    try:
        request = json.loads(sys.stdin.read(25_000_000))
        body = base64.b64decode(request['body'], validate=True)
        if len(body) > MAX_BYTES:
            raise SourceError('unsupported', 'import', 'import_size_limit')
        with contextlib.redirect_stdout(sys.stderr):
            parsed = {'docx': parse_docx, 'xlsx': parse_xlsx, 'csv': parse_csv,
                      'image': parse_image, 'scan_pdf': parse_scan_pdf}[request['kind']](body, request['title'])
        print(json.dumps(parsed, ensure_ascii=False))
    except SourceError as error:
        print(json.dumps({'error': {'status': error.status, 'stage': error.stage, 'reason': error.reason}}))
        sys.exit(1)
    except Exception:
        print(json.dumps({'error': {'status': 'parse_error', 'stage': 'import', 'reason': 'invalid_local_document'}}))
        sys.exit(1)
